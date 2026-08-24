"""Reading and writing the Design Project Close-Out workbook.

The close-out data has always lived in ``DesignProjectCloseOutAssetTable_master.xlsx``,
and the risk ETL still reads that file. Both directions here therefore treat the
workbook's ``Template`` sheet as a contract: twelve columns, exact headings, one row per
asset. The export is a drop-in replacement for the master file; the import accepts both
a single-project template and the full master, grouping rows into projects by the four
attributes that identify one analysis batch.
"""

from __future__ import annotations

import re
from datetime import date, datetime, timedelta
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook

TEMPLATE_SHEET = "Template"
TEMPLATE_HEADINGS = (
    "Asset ID",
    "Construction Plan ID",
    "Critical Facility ID",
    "Source of Analysis",
    "Cityworks WO ID",
    "Project Name",
    "Date of Analysis",
    "Flooding Design Standards",
    "Flooding Impact",
    "Flooding Service Eligibility",
    "Post Project Asset Condition",
    "Notes",
)

ASSET_ID_PATTERN = re.compile(r"^[PSD]_[A-Za-z0-9]+$")
# The master template embeds one instruction row whose cells describe each column
# ("Pipe (P_######), ..."). Those rows are boilerplate, not data.
_INSTRUCTION_MARKER = "######"

IMPORT_ROW_LIMIT = 20_000


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def normalize_analysis_date(value: Any) -> str | None:
    """Return YYYY-MM-DD, accepting the datetime cells and text variants Excel holds."""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    text = str(value or "").strip()
    if not text:
        return None
    for pattern in (
        "%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%Y/%m/%d",
        "%m/%d/%Y", "%m-%d-%Y", "%m/%d/%y", "%m-%d-%y",
        "%d-%b-%Y", "%d %b %Y", "%b %d, %Y", "%B %d, %Y",
    ):
        try:
            return datetime.strptime(text[:19], pattern).strftime("%Y-%m-%d")
        except ValueError:
            continue
    swapped = re.fullmatch(r"(\d{4})[-/](\d{1,2})[-/](\d{1,2})", text)
    if swapped:
        year, middle, last = (int(part) for part in swapped.groups())
        # Only when the middle number cannot be a month: "2024-23-01" has one reading,
        # 23 January. Anything genuinely ambiguous is left for a person to fix.
        if middle > 12 and 1 <= last <= 12:
            try:
                return date(year, last, middle).strftime("%Y-%m-%d")
            except ValueError:
                return None
    if text.isdigit() and 20000 <= int(text) <= 60000:
        # A date column formatted as text keeps Excel's serial number instead of a date.
        return (date(1899, 12, 30) + timedelta(days=int(text))).strftime("%Y-%m-%d")
    return None


def parse_template_workbook(content: bytes) -> tuple[list[dict[str, str]], list[str]]:
    """Return (rows, errors) from a close-out workbook's Template sheet.

    Rows are keyed by the template headings, values as trimmed text (dates normalised to
    YYYY-MM-DD). Structural problems abort with a single error; per-row problems are
    reported with their worksheet row number and exclude only the offending row.
    """
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception:
        return [], ["The file could not be read as an Excel workbook (.xlsx)."]
    try:
        if TEMPLATE_SHEET in workbook.sheetnames:
            sheet = workbook[TEMPLATE_SHEET]
        else:
            sheet = workbook[workbook.sheetnames[0]]

        rows_iterator = sheet.iter_rows(values_only=True)
        header = next(rows_iterator, None)
        if header is None:
            return [], [f"The {sheet.title} sheet is empty."]
        headings = [_cell_text(value) for value in header[: len(TEMPLATE_HEADINGS)]]
        if tuple(headings) != TEMPLATE_HEADINGS:
            return [], [
                "The sheet does not match the close-out template. Expected columns: "
                + ", ".join(TEMPLATE_HEADINGS)
            ]

        rows: list[dict[str, str]] = []
        errors: list[str] = []
        for index, values in enumerate(rows_iterator, start=2):
            if index - 1 > IMPORT_ROW_LIMIT:
                errors.append(f"The import stops after {IMPORT_ROW_LIMIT:,} rows.")
                break
            cells = list(values[: len(TEMPLATE_HEADINGS)])
            cells += [None] * (len(TEMPLATE_HEADINGS) - len(cells))
            if all(_cell_text(value) == "" for value in cells):
                continue
            if any(_INSTRUCTION_MARKER in _cell_text(value) for value in cells):
                continue
            row = dict(zip(TEMPLATE_HEADINGS, (_cell_text(value) for value in cells)))
            analysis_date = normalize_analysis_date(cells[6])
            if analysis_date is not None:
                row["Date of Analysis"] = analysis_date
            elif row["Date of Analysis"]:
                errors.append(
                    f"Row {index}: Date of Analysis {row['Date of Analysis']!r} is not a date."
                )
                continue
            rows.append(row)
        return rows, errors
    finally:
        workbook.close()


def group_rows_into_projects(rows: list[dict[str, str]]) -> list[dict[str, Any]]:
    """Group template rows into projects by the four batch-identity attributes.

    A single-project template collapses to one group; the historical master file
    yields one project per (name, WO, source, date) combination, preserving even the
    inconsistent legacy combinations as-is rather than silently repairing them.
    """
    groups: dict[tuple[str, str, str, str], dict[str, Any]] = {}
    for row in rows:
        key = (
            row["Project Name"],
            row["Cityworks WO ID"],
            row["Source of Analysis"],
            row["Date of Analysis"],
        )
        group = groups.setdefault(
            key,
            {
                "project_name": row["Project Name"] or None,
                "cityworks_wo_id": row["Cityworks WO ID"] or None,
                "source_of_analysis": row["Source of Analysis"] or None,
                "date_of_analysis": row["Date of Analysis"] or None,
                "assets": [],
            },
        )
        group["assets"].append(
            {
                "asset_id": row["Asset ID"],
                "construction_plan_id": row["Construction Plan ID"] or None,
                "critical_facility_id": row["Critical Facility ID"] or None,
                "flooding_design_standards": row["Flooding Design Standards"] or None,
                "flooding_impact": row["Flooding Impact"] or None,
                "flooding_service_eligibility": row["Flooding Service Eligibility"] or None,
                "post_project_asset_condition": row["Post Project Asset Condition"] or None,
                "notes": row["Notes"] or None,
            }
        )
    return list(groups.values())


def build_template_workbook(rows: list[dict[str, Any]]) -> bytes:
    """Build a plain Template-sheet workbook the risk ETL can read as the master file."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = TEMPLATE_SHEET
    sheet.append(list(TEMPLATE_HEADINGS))
    for row in rows:
        sheet.append([row.get(heading) for heading in TEMPLATE_HEADINGS])
    for column, width in zip("ABCDEFGHIJKL", (14, 18, 16, 18, 14, 28, 14, 24, 24, 34, 26, 40)):
        sheet.column_dimensions[column].width = width
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
