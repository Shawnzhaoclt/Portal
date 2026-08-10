from __future__ import annotations

import re
from collections.abc import Callable, Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime, time
from io import BytesIO
from typing import Any, Literal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins, PrintPageSetup
from openpyxl.worksheet.properties import PageSetupProperties


ExcelDataType = Literal["text", "number", "integer", "date", "datetime"]
ExcelTransform = Callable[[Any], Any]

PORTAL_NAVY = "0B3D68"
PORTAL_BLUE = "1F5D8F"
PORTAL_LIGHT_BLUE = "D6E7F2"
PORTAL_PALE_BLUE = "EEF5FA"
PORTAL_ALTERNATE_ROW = "EAF3FA"
PORTAL_BORDER = "B7C9D8"
PORTAL_TEXT = "183247"
PORTAL_MUTED_TEXT = "536B7E"
PORTAL_WHITE = "FFFFFF"
HEADER_ROW = 4


@dataclass(frozen=True, slots=True)
class ExcelColumn:
    key: str
    heading: str
    width: float = 18
    data_type: ExcelDataType = "text"
    wrap_text: bool = False
    transform: ExcelTransform | None = None


def _safe_sheet_name(value: str) -> str:
    normalized = re.sub(r"[\\/*?:\[\]]", " ", str(value or "Export")).strip()
    return (normalized or "Export")[:31]


def _local_naive_datetime(value: Any) -> datetime | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, time.min)
    else:
        normalized = str(value).strip().replace("Z", "+00:00")
        try:
            parsed = datetime.fromisoformat(normalized)
        except ValueError:
            return None
    if parsed.tzinfo is not None:
        parsed = parsed.astimezone().replace(tzinfo=None)
    return parsed.replace(microsecond=0)


def _safe_text(value: Any) -> str | None:
    if value in (None, ""):
        return None
    text_value = str(value)
    # Prevent user-entered text from becoming a spreadsheet formula.
    return f"'{text_value}" if text_value.startswith(("=", "+", "-", "@")) else text_value


def _cell_value(column: ExcelColumn, row: Mapping[str, Any]) -> Any:
    value = row.get(column.key)
    if column.transform is not None:
        value = column.transform(value)
    if value in (None, ""):
        return None
    if column.data_type == "datetime":
        return _local_naive_datetime(value) or _safe_text(value)
    if column.data_type == "date":
        parsed = _local_naive_datetime(value)
        return parsed.date() if parsed else _safe_text(value)
    if column.data_type == "integer":
        try:
            return int(value)
        except (TypeError, ValueError):
            return _safe_text(value)
    if column.data_type == "number":
        try:
            return float(value)
        except (TypeError, ValueError):
            return _safe_text(value)
    return _safe_text(value)


def _metadata_filters(filters: Mapping[str, Any] | None) -> str:
    parts = [
        f"{label}: {value}"
        for label, value in (filters or {}).items()
        if value not in (None, "")
    ]
    return "Filters: " + ("; ".join(parts) if parts else "All records")


def _style_merged_row(
    sheet: Any,
    row: int,
    column_count: int,
    *,
    value: str,
    fill: str,
    color: str,
    size: int,
    bold: bool,
    height: float,
) -> None:
    for column_index in range(1, column_count + 1):
        cell = sheet.cell(row=row, column=column_index)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.border = Border(bottom=Side(style="thin", color=PORTAL_BORDER))
    if column_count > 1:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=column_count)
    anchor = sheet.cell(row=row, column=1, value=value)
    anchor.font = Font(name="Aptos", size=size, bold=bold, color=color)
    anchor.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[row].height = height


def build_portal_excel_export(
    *,
    report_title: str,
    columns: Sequence[ExcelColumn],
    rows: Sequence[Mapping[str, Any]],
    exported_by: str,
    filters: Mapping[str, Any] | None = None,
    sheet_name: str | None = None,
    generated_at: datetime | None = None,
) -> bytes:
    """Build a consistently styled, filterable Portal Excel workbook."""

    if not columns:
        raise ValueError("At least one Excel export column is required.")

    generated = (generated_at or datetime.now().astimezone()).replace(microsecond=0)
    row_count = len(rows)
    column_count = len(columns)
    workbook = Workbook()
    workbook.properties.creator = "City of Charlotte Storm Water Services"
    workbook.properties.title = report_title
    workbook.properties.subject = "Portal data export"
    workbook.properties.description = f"{report_title} export containing {row_count} records."

    sheet = workbook.active
    sheet.title = _safe_sheet_name(sheet_name or report_title)
    sheet.sheet_properties.tabColor = PORTAL_BLUE
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True, autoPageBreaks=False)
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 90

    _style_merged_row(
        sheet,
        1,
        column_count,
        value=report_title,
        fill=PORTAL_BLUE,
        color=PORTAL_WHITE,
        size=16,
        bold=True,
        height=30,
    )
    record_label = "record" if row_count == 1 else "records"
    _style_merged_row(
        sheet,
        2,
        column_count,
        value=f"Generated: {generated.strftime('%Y-%m-%d %H:%M:%S')} | Exported by: {exported_by} | {row_count:,} {record_label}",
        fill=PORTAL_WHITE,
        color=PORTAL_MUTED_TEXT,
        size=10,
        bold=False,
        height=22,
    )
    _style_merged_row(
        sheet,
        3,
        column_count,
        value=_metadata_filters(filters),
        fill=PORTAL_PALE_BLUE,
        color=PORTAL_MUTED_TEXT,
        size=9,
        bold=False,
        height=28,
    )

    grid_side = Side(style="thin", color=PORTAL_BORDER)
    grid_border = Border(left=grid_side, right=grid_side, top=grid_side, bottom=grid_side)
    for column_index, column in enumerate(columns, start=1):
        cell = sheet.cell(row=HEADER_ROW, column=column_index, value=column.heading)
        cell.font = Font(name="Aptos", size=10, bold=True, color=PORTAL_NAVY)
        cell.fill = PatternFill("solid", fgColor=PORTAL_LIGHT_BLUE)
        cell.border = grid_border
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.column_dimensions[cell.column_letter].width = max(8, min(float(column.width), 60))
    sheet.row_dimensions[HEADER_ROW].height = 30

    for row_index, record in enumerate(rows, start=HEADER_ROW + 1):
        alternate_fill = PatternFill("solid", fgColor=PORTAL_ALTERNATE_ROW if row_index % 2 == 1 else PORTAL_WHITE)
        for column_index, column in enumerate(columns, start=1):
            cell = sheet.cell(row=row_index, column=column_index, value=_cell_value(column, record))
            cell.font = Font(name="Aptos", size=10, color=PORTAL_TEXT)
            cell.fill = alternate_fill
            cell.border = grid_border
            cell.alignment = Alignment(vertical="top", wrap_text=column.wrap_text)
            if column.data_type == "datetime" and isinstance(cell.value, datetime):
                cell.number_format = "m/d/yyyy h:mm:ss AM/PM"
            elif column.data_type == "date" and isinstance(cell.value, date):
                cell.number_format = "m/d/yyyy"
            elif column.data_type == "integer":
                cell.number_format = "0"
            elif column.data_type == "number":
                cell.number_format = "0.0"
        sheet.row_dimensions[row_index].height = 30 if any(column.wrap_text for column in columns) else 20

    last_column = sheet.cell(row=HEADER_ROW, column=column_count).column_letter
    last_row = max(HEADER_ROW, HEADER_ROW + row_count)
    sheet.freeze_panes = f"A{HEADER_ROW + 1}"
    sheet.auto_filter.ref = f"A{HEADER_ROW}:{last_column}{last_row}"
    sheet.print_area = f"A1:{last_column}{last_row}"
    sheet.print_title_rows = f"1:{HEADER_ROW}"
    sheet.page_setup = PrintPageSetup(orientation="landscape", paperSize="9", fitToWidth=1, fitToHeight=0)
    sheet.page_margins = PageMargins(left=0.25, right=0.25, top=0.5, bottom=0.5, header=0.2, footer=0.2)
    sheet.oddFooter.center.text = "Page &P of &N"
    sheet.oddFooter.center.size = 9
    sheet.oddFooter.center.color = PORTAL_MUTED_TEXT

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
